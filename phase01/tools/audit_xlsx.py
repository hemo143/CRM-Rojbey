"""Offline, read-only workbook audit. Never prints customer values or writes into Git."""
import argparse, collections, datetime, hashlib, json, re
from pathlib import Path
import openpyxl

def phone(value):
    value=str(value or '').strip().translate(str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹','01234567890123456789'))
    value=re.sub(r'[\s().-]','',value)
    return '+'+value[2:] if value.startswith('00') else value

def encode(value):
    if isinstance(value,(datetime.datetime,datetime.date)):
        return {'__xlsx_date':value.isoformat()}
    return value

def audit(path,snapshot=None):
    original=path.read_bytes()
    book=openpyxl.load_workbook(path,data_only=False)
    sheets=[]
    for s in book:
        occupied=[r for r in s.iter_rows() if any(c.value is not None for c in r)]
        sheets.append({'name':s.title,'dimensions':[s.max_row,s.max_column],'nonempty_data_rows':max(0,len(occupied)-1),'headers':[c.value for c in s[1]],'formulas':sum(c.data_type=='f' for r in occupied for c in r),'validation_rules':[{'range':str(v.sqref),'type':v.type,'source_type':'inline' if str(v.formula1).startswith('"') else 'range_or_formula'} for v in s.data_validations.dataValidation],'freeze':str(s.freeze_panes),'filter':s.auto_filter.ref,'rtl':s.sheet_view.rightToLeft})
    s=book['العملاء'];heads=[c.value for c in s[1]]
    rows=[dict(zip(heads,r)) for r in s.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
    ids=collections.Counter(r['Customer ID'] for r in rows if r['Customer ID'])
    phones=collections.Counter(phone(r['رقم الهاتف']) for r in rows if r['رقم الهاتف'])
    quality={'customers':len(rows),'missing_ids':sum(not r['Customer ID'] for r in rows),'duplicate_id_groups':sum(v>1 for v in ids.values()),'missing_names':sum(not r['الاسم'] for r in rows),'missing_phones':sum(not r['رقم الهاتف'] for r in rows),'numeric_phones':sum(isinstance(r['رقم الهاتف'],(int,float)) for r in rows),'duplicate_phone_groups':sum(v>1 for v in phones.values()),'records_in_duplicate_phone_groups':sum(v for v in phones.values() if v>1),'invalid_phone_shapes':sum(not bool(re.fullmatch(r'\+?\d{7,15}',v)) for v in phones),'missing_sources':sum(not r['مصدر العميلة'] for r in rows),'source_distinct_count':len({r['مصدر العميلة'] for r in rows if r['مصدر العميلة']}),'historical_staff_distinct_count':len({r['الموظفة'] for r in rows if r['الموظفة']}),'blank_created_at':sum(not r['تاريخ الإضافة'] for r in rows),'blank_updated_at':sum(not r['آخر تعديل'] for r in rows),'status_counts':dict(collections.Counter(str(r['حالة العميلة'] or '(blank)') for r in rows))}
    result={'scope':'uploaded XLSX snapshot only; not live Google','sha256':hashlib.sha256(original).hexdigest(),'sheets':sheets,'named_range_count':len(book.defined_names),'quality':quality}
    if snapshot:
        snapshot=snapshot.resolve();repo=Path(__file__).resolve().parents[2]
        if snapshot.is_relative_to(repo):raise ValueError('Private snapshot must remain outside the repository')
        snapshot.parent.mkdir(parents=True,exist_ok=True,mode=0o700)
        payload={'sheets':[{'name':s.title,'rows':[[encode(c.value) for c in r] for r in s.iter_rows()]} for s in book]}
        # Open with owner-only permissions before writing any customer values.
        import os
        fd=os.open(snapshot,os.O_WRONLY|os.O_CREAT|os.O_TRUNC,0o600)
        os.fchmod(fd,0o600)
        with os.fdopen(fd,'w') as output:json.dump(payload,output,ensure_ascii=False)
        snapshot.chmod(0o600)
        result['private_snapshot_written']=True
    if path.read_bytes()!=original:raise RuntimeError('Source workbook changed unexpectedly')
    result['source_unchanged']=True
    return result
if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('workbook',type=Path);parser.add_argument('--snapshot-private',type=Path)
    args=parser.parse_args();print(json.dumps(audit(args.workbook,args.snapshot_private),ensure_ascii=False,indent=2))
